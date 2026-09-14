import csv
import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.models import MarketSummary, StockQuote

LOGGER = logging.getLogger(__name__)


class ExcelExporter:
    """Schreibt die verarbeiteten Boersendaten in eine formatierte Excel-Datei."""

    HEADERS = [
        "Symbol",
        "Name",
        "Waehrung",
        "Aktueller Kurs",
        "Vortagesschluss",
        "Veraenderung",
        "Veraenderung %",
        "Tageshoch",
        "Tagestief",
        "Handelsspanne %",
        "Handelsvolumen",
        "Datenzeit",
    ]

    def export(self, quotes: list[StockQuote], summary: MarketSummary, output_path: str) -> str:
        if not quotes:
            raise ValueError("Keine Daten fuer den Excel-Export vorhanden.")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Kursdaten"
        sheet.freeze_panes = "A5"

        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        info_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        thin = Side(style="thin", color="B7C9D6")
        border = Border(bottom=thin)
        last_column = "L"

        sheet.merge_cells(f"A1:{last_column}1")
        sheet["A1"] = "Kursuebersicht mit Tagesveraenderung"
        sheet["A1"].font = Font(size=14, bold=True)
        sheet["A1"].alignment = Alignment(horizontal="center")

        sheet.merge_cells(f"A2:{last_column}2")
        sheet["A2"] = (
            f"Stand: {quotes[0].fetched_at:%d.%m.%Y %H:%M:%S}  |  "
            f"{summary.count} Wertpapiere  |  {summary.gainers} im Plus / {summary.losers} im Minus"
        )
        sheet["A2"].fill = info_fill
        sheet["A2"].alignment = Alignment(horizontal="left")

        header_row = 4
        for col_index, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row=header_row, column=col_index, value=header)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        for row_index, quote in enumerate(quotes, start=5):
            sheet.cell(row=row_index, column=1, value=quote.symbol)
            sheet.cell(row=row_index, column=2, value=quote.name)
            sheet.cell(row=row_index, column=3, value=quote.currency)
            sheet.cell(row=row_index, column=4, value=quote.current_price)
            sheet.cell(row=row_index, column=5, value=quote.previous_close)
            # Formeln statt fester Werte: die Datei bleibt in Excel nachvollziehbar.
            sheet.cell(row=row_index, column=6, value=f"=D{row_index}-E{row_index}")
            sheet.cell(row=row_index, column=7, value=f"=IF(E{row_index}=0,0,F{row_index}/E{row_index})")
            sheet.cell(row=row_index, column=8, value=quote.day_high)
            sheet.cell(row=row_index, column=9, value=quote.day_low)
            sheet.cell(row=row_index, column=10, value=f"=IF(I{row_index}=0,0,(H{row_index}-I{row_index})/I{row_index})")
            sheet.cell(row=row_index, column=11, value=quote.volume)
            sheet.cell(row=row_index, column=12, value=quote.fetched_at)

        data_start = 5
        data_end = 4 + len(quotes)

        self._apply_number_format(sheet, data_start, data_end, columns=(4, 6), number_format="#,##0.00")
        self._apply_number_format(sheet, data_start, data_end, columns=(5, 5), number_format="#,##0.00")
        self._apply_number_format(sheet, data_start, data_end, columns=(7, 7), number_format="0.00%")
        self._apply_number_format(sheet, data_start, data_end, columns=(8, 9), number_format="#,##0.00")
        self._apply_number_format(sheet, data_start, data_end, columns=(10, 10), number_format="0.00%")
        self._apply_number_format(sheet, data_start, data_end, columns=(11, 11), number_format="#,##0")
        self._apply_number_format(sheet, data_start, data_end, columns=(12, 12), number_format="yyyy-mm-dd hh:mm:ss")

        widths = {
            "A": 12, "B": 26, "C": 10, "D": 15, "E": 16, "F": 14,
            "G": 15, "H": 13, "I": 13, "J": 16, "K": 17, "L": 21,
        }
        for column_letter, width in widths.items():
            sheet.column_dimensions[column_letter].width = width

        # Farbskala fuer die Tagesveraenderung: rot = Verlust, gruen = Gewinn.
        sheet.conditional_formatting.add(
            f"G{data_start}:G{data_end}",
            ColorScaleRule(
                start_type="min", start_color="F4B183",
                mid_type="num", mid_value=0, mid_color="FFFFFF",
                end_type="max", end_color="A9D18E",
            ),
        )
        # Datenbalken fuer die Handelsspanne: Volatilitaet auf einen Blick.
        sheet.conditional_formatting.add(
            f"J{data_start}:J{data_end}",
            DataBarRule(start_type="min", end_type="max", color="4F81BD", showValue=True),
        )

        self._add_summary_sheet(workbook, summary)
        self._add_chart_sheet(workbook, quotes)

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
        LOGGER.info("Excel-Datei geschrieben: %s", output)
        return str(output)

    @staticmethod
    def _apply_number_format(sheet, start: int, end: int, columns: tuple[int, int], number_format: str) -> None:
        for row in sheet.iter_rows(min_row=start, max_row=end, min_col=columns[0], max_col=columns[1]):
            for cell in row:
                cell.number_format = number_format
                cell.alignment = Alignment(vertical="center")

    @staticmethod
    def _add_summary_sheet(workbook: Workbook, summary: MarketSummary) -> None:
        sheet = workbook.create_sheet("Kennzahlen")
        sheet["A1"] = "Auswertung des Abrufs"
        sheet["A1"].font = Font(size=14, bold=True)

        rows = [
            ("Anzahl Wertpapiere", summary.count),
            ("Im Plus", summary.gainers),
            ("Im Minus", summary.losers),
            ("Unveraendert", summary.unchanged),
            ("Durchschnittliche Veraenderung in %", summary.average_change_percent),
            ("Gesamtes Handelsvolumen", summary.total_volume),
            ("Tagesbester", f"{summary.best.symbol} ({summary.best.change_percent} %)" if summary.best else "-"),
            ("Tagesschlechtester", f"{summary.worst.symbol} ({summary.worst.change_percent} %)" if summary.worst else "-"),
            (
                "Hoechste Handelsspanne",
                f"{summary.most_volatile.symbol} ({summary.most_volatile.day_range_percent} %)"
                if summary.most_volatile
                else "-",
            ),
        ]
        for row_index, (label, value) in enumerate(rows, start=3):
            sheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
            sheet.cell(row=row_index, column=2, value=value)

        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 28

    @staticmethod
    def _add_chart_sheet(workbook: Workbook, quotes: list[StockQuote]) -> None:
        sheet = workbook.create_sheet("Visualisierung")
        sheet["A1"] = "Tagesveraenderung je Symbol"
        sheet["A1"].font = Font(size=14, bold=True)
        sheet["A3"] = "Symbol"
        sheet["B3"] = "Veraenderung %"
        sheet["A3"].font = sheet["B3"].font = Font(bold=True)

        for row_index, quote in enumerate(quotes, start=4):
            sheet.cell(row=row_index, column=1, value=quote.symbol)
            sheet.cell(row=row_index, column=2, value=quote.change_percent)

        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Veraenderung gegenueber Vortag in %"
        chart.y_axis.title = "Symbol"
        chart.x_axis.title = "Prozent"
        data = Reference(sheet, min_col=2, min_row=3, max_row=3 + len(quotes))
        categories = Reference(sheet, min_col=1, min_row=4, max_row=3 + len(quotes))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 9
        chart.width = 17
        sheet.add_chart(chart, "D3")
        sheet.column_dimensions["A"].width = 14
        sheet.column_dimensions["B"].width = 16


class CsvExporter:
    """Schreibt die Kursdaten als CSV-Datei (Semikolon, Excel-tauglich)."""

    HEADERS = [
        "Symbol", "Name", "Waehrung", "Kurs", "Vortagesschluss", "Veraenderung",
        "Veraenderung_Prozent", "Tageshoch", "Tagestief", "Handelsspanne_Prozent",
        "Volumen", "Datenzeit",
    ]

    def export(self, quotes: list[StockQuote], output_path: str) -> str:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(self.HEADERS)
            for quote in quotes:
                writer.writerow([
                    quote.symbol,
                    quote.name,
                    quote.currency,
                    quote.current_price,
                    quote.previous_close,
                    quote.change_absolute,
                    quote.change_percent,
                    quote.day_high,
                    quote.day_low,
                    quote.day_range_percent,
                    quote.volume,
                    quote.fetched_at.isoformat(timespec="seconds"),
                ])

        LOGGER.info("CSV-Datei geschrieben: %s", output)
        return str(output)


class JsonExporter:
    """Schreibt Kursdaten und Kennzahlen als JSON fuer die Weiterverarbeitung."""

    def export(self, quotes: list[StockQuote], summary: MarketSummary, output_path: str) -> str:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        payload = {
            "abgerufen_am": quotes[0].fetched_at.isoformat(timespec="seconds") if quotes else None,
            "kennzahlen": {
                "anzahl": summary.count,
                "im_plus": summary.gainers,
                "im_minus": summary.losers,
                "unveraendert": summary.unchanged,
                "durchschnittliche_veraenderung_prozent": summary.average_change_percent,
                "gesamtvolumen": summary.total_volume,
                "tagesbester": summary.best.symbol if summary.best else None,
                "tagesschlechtester": summary.worst.symbol if summary.worst else None,
                "hoechste_handelsspanne": summary.most_volatile.symbol if summary.most_volatile else None,
            },
            "kursdaten": [
                {
                    "symbol": quote.symbol,
                    "name": quote.name,
                    "waehrung": quote.currency,
                    "kurs": quote.current_price,
                    "vortagesschluss": quote.previous_close,
                    "veraenderung": quote.change_absolute,
                    "veraenderung_prozent": quote.change_percent,
                    "tageshoch": quote.day_high,
                    "tagestief": quote.day_low,
                    "handelsspanne_prozent": quote.day_range_percent,
                    "volumen": quote.volume,
                    "datenzeit": quote.fetched_at.isoformat(timespec="seconds"),
                }
                for quote in quotes
            ],
        }

        output.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        LOGGER.info("JSON-Datei geschrieben: %s", output)
        return str(output)
